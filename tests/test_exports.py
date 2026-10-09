"""XLSX export: localized headers and values, and the ``lang`` query param."""

from __future__ import annotations

import io

import httpx
from openpyxl import load_workbook

from driftwatch.api.exports import _significance


async def _sheet(client: httpx.AsyncClient, url: str) -> tuple[str, list[str]]:
    response = await client.get(url)
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    assert sheet is not None
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    return sheet.title, headers


async def test_export_defaults_to_english(admin_client: httpx.AsyncClient) -> None:
    title, headers = await _sheet(admin_client, "/api/exports/changes.xlsx")
    assert title == "Changes"
    assert headers == [
        "Change",
        "Site",
        "URL",
        "Detected (UTC)",
        "Significant",
        "Headline",
        "Summary",
        "Notified",
    ]


async def test_export_renders_polish_headers(admin_client: httpx.AsyncClient) -> None:
    title, headers = await _sheet(admin_client, "/api/exports/changes.xlsx?lang=pl")
    assert title == "Zmiany"
    assert headers == [
        "Zmiana",
        "Strona",
        "URL",
        "Wykryto (UTC)",
        "Istotna",
        "Nagłówek",
        "Podsumowanie",
        "Powiadomiono",
    ]


async def test_export_unknown_language_falls_back_to_english(
    admin_client: httpx.AsyncClient,
) -> None:
    title, _ = await _sheet(admin_client, "/api/exports/changes.xlsx?lang=de")
    assert title == "Changes"


def test_significance_values_are_localized() -> None:
    assert _significance(None, "pl") == "oczekuje"
    assert _significance(True, "pl") == "tak"
    assert _significance(False, "pl") == "nie"
    assert _significance(None, "en") == "pending"
