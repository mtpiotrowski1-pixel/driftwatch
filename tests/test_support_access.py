"""Public tenant support access requires a scoped break-glass grant."""

from __future__ import annotations

import io
from collections.abc import AsyncIterator
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from pathlib import Path

import httpx
import pytest
import pytest_asyncio
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import select, text, update

from driftwatch.api.deps import STEP_UP_COOKIE
from driftwatch.app import create_app
from driftwatch.config import Settings
from driftwatch.db import Database
from driftwatch.enums import (
    CheckJobKind,
    CheckJobSource,
    CheckJobStatus,
    NotificationStatus,
)
from driftwatch.models import (
    AIUsage,
    AuditEvent,
    ChangeEvent,
    NotificationLog,
    Site,
    SiteCheckJob,
    Snapshot,
    SupportAccessGrant,
)
from driftwatch.monitoring.analyzer import Analysis
from driftwatch.monitoring.pipeline import CheckStatus
from driftwatch.monitoring.usage import CostEstimate, TokenUsage
from driftwatch.runner import RunResult
from driftwatch.security import two_factor
from driftwatch.security.tokens import issue_support_access_token
from tests.conftest import FakePicker, RecordingChannel, ScriptedCapturer, StubAnalyzer

_OPERATOR_EMAIL = "operator@example.com"
_OPERATOR_PASSWORD = "public-operator-password"
_TEST_SESSION_GENERATION = "2f8d2a36-9de4-4b61-9f0e-0b5fe35b19cf"


@dataclass(frozen=True, slots=True)
class PublicOperatorSession:
    client: httpx.AsyncClient
    replay_client: httpx.AsyncClient
    database: Database
    totp_secret: str
    user_id: int
    app: FastAPI


@pytest_asyncio.fixture
async def public_operator_session(tmp_path: Path) -> AsyncIterator[PublicOperatorSession]:
    settings = Settings(
        data_dir=tmp_path,
        database_url=f"sqlite+aiosqlite:///{tmp_path}/support.db",
        base_url="https://driftwatch.example.test",
        host="0.0.0.0",
        session_secret_key="session-key-for-public-support-tests-0123456789",
        encryption_key="encryption-key-for-public-support-tests-012345",
        initial_admin_email=_OPERATOR_EMAIL,
        initial_admin_password=_OPERATOR_PASSWORD,
        scheduler_enabled=False,
        picker_enabled=False,
        run_migrations=True,
    )
    app = create_app(
        settings,
        capturer=ScriptedCapturer(html="<main>ok</main>"),
        analyzer=StubAnalyzer(),
        channel=RecordingChannel(),
        picker=FakePicker(),
        enable_scheduler=False,
    )
    async with app.router.lifespan_context(app):
        transport = httpx.ASGITransport(app=app)
        async with (
            httpx.AsyncClient(
                transport=transport,
                base_url=settings.base_url,
            ) as client,
            httpx.AsyncClient(
                transport=httpx.ASGITransport(app=app),
                base_url=settings.base_url,
            ) as replay_client,
        ):
            login = await client.post(
                "/api/auth/login",
                json={"email": _OPERATOR_EMAIL, "password": _OPERATOR_PASSWORD},
            )
            assert login.status_code == 200, login.text
            enrollment_step_up = await client.post(
                "/api/auth/step-up",
                json={"password": _OPERATOR_PASSWORD},
            )
            assert enrollment_step_up.status_code == 204, enrollment_step_up.text
            setup = await client.post("/api/auth/totp/setup")
            assert setup.status_code == 200, setup.text
            secret = setup.json()["secret"]
            enable = await client.post(
                "/api/auth/totp/enable",
                json={"code": two_factor.generate_code(secret)},
            )
            assert enable.status_code == 200, enable.text
            yield PublicOperatorSession(
                client,
                replay_client,
                app.state.db,
                secret,
                int(login.json()["user"]["id"]),
                app,
            )


async def _create_organization(session: PublicOperatorSession, name: str) -> int:
    await _step_up(session)
    response = await session.client.post("/api/organizations", json={"name": name})
    assert response.status_code == 201, response.text
    session.client.cookies.delete(STEP_UP_COOKIE)
    return int(response.json()["id"])


async def _step_up(
    session: PublicOperatorSession,
    organization_id: int | None = None,
) -> None:
    client = session.client
    headers = {"X-Acting-Org": str(organization_id)} if organization_id is not None else None
    response = await client.post(
        "/api/auth/step-up",
        headers=headers,
        json={
            "password": _OPERATOR_PASSWORD,
            "totp_code": two_factor.generate_code(session.totp_secret),
        },
    )
    assert response.status_code == 204, response.text


async def _grant(session: PublicOperatorSession, organization_id: int) -> dict[str, object]:
    client = session.client
    await _step_up(session, organization_id)
    response = await client.post(
        "/api/support-access",
        headers={"X-Acting-Org": str(organization_id)},
        json={"reason": "Investigating customer incident", "ticket": "SUP-123"},
    )
    assert response.status_code == 200, response.text
    return response.json()


async def _seed_tenant_history(
    session: PublicOperatorSession,
    organization_id: int,
    *,
    marker: str,
) -> None:
    async with session.database.session() as database_session:
        site = Site(
            organization_id=organization_id,
            url=f"https://{marker}.example.test",
            name=f"Site {marker}",
        )
        database_session.add(site)
        await database_session.flush()
        snapshot = Snapshot(
            site_id=site.id, content_html="<main>changed</main>", content_text="changed"
        )
        database_session.add(snapshot)
        await database_session.flush()
        change = ChangeEvent(
            site_id=site.id,
            new_snapshot_id=snapshot.id,
            significant=True,
            user_verdict=True,
            headline=f"Headline {marker}",
            summary=f"Summary {marker}",
        )
        database_session.add(change)
        await database_session.flush()
        database_session.add_all(
            [
                NotificationLog(
                    change_id=change.id,
                    recipient_email=f"{marker}@example.test",
                    channel="email",
                    status=NotificationStatus.SENT,
                ),
                AIUsage(
                    organization_id=organization_id,
                    site_id=site.id,
                    change_id=change.id,
                    model=f"model-{marker}",
                    prompt_tokens=10,
                    completion_tokens=5,
                    total_tokens=15,
                    cost_usd=0.01,
                ),
            ]
        )
        await database_session.commit()


@pytest.mark.parametrize("action", ["check", "snapshot", "retry", "analyze", "analyze-preview"])
async def test_supported_work_releases_audit_transaction_before_delegate(
    public_operator_session: PublicOperatorSession,
    monkeypatch: pytest.MonkeyPatch,
    action: str,
) -> None:
    operator = public_operator_session
    organization_id = await _create_organization(operator, "Delegated support customer")
    other_organization_id = await _create_organization(operator, "Other support customer")
    await _seed_tenant_history(operator, organization_id, marker="delegated")
    await _seed_tenant_history(operator, other_organization_id, marker="other-delegated")
    async with operator.database.session() as database_session:
        own_site = await database_session.scalar(
            select(Site).where(Site.organization_id == organization_id)
        )
        other_site = await database_session.scalar(
            select(Site).where(Site.organization_id == other_organization_id)
        )
        assert own_site is not None and other_site is not None
        own_change = await database_session.scalar(
            select(ChangeEvent).where(ChangeEvent.site_id == own_site.id)
        )
        other_change = await database_session.scalar(
            select(ChangeEvent).where(ChangeEvent.site_id == other_site.id)
        )
        assert own_change is not None and other_change is not None
        own_site_id, own_change_id = own_site.id, own_change.id
        other_site_id, other_change_id = other_site.id, other_change.id

    captures_page = action in {"check", "snapshot"}
    resource_type = "sites" if captures_page else "changes"
    resource_id = own_site_id if captures_page else own_change_id
    other_resource_id = other_site_id if captures_page else other_change_id
    path = f"/api/{resource_type}/{resource_id}/{action}"
    other_path = f"/api/{resource_type}/{other_resource_id}/{action}"
    headers = {"X-Acting-Org": str(organization_id)}
    payload = {"rules": "Only important changes"} if action == "analyze-preview" else None
    calls: list[int] = []

    async def delegated_work(
        delegated_id: int, *args: object, **kwargs: object
    ) -> Analysis | RunResult:
        calls.append(delegated_id)
        assert delegated_id == resource_id
        expected_args = ("Only important changes",) if action == "analyze-preview" else ()
        expected_kwargs = {"analyze": True} if action == "check" else {}
        if action == "retry":
            expected_kwargs = {"force_delivery": True}
        assert args == expected_args
        assert kwargs == expected_kwargs

        # A separate engine exercises a real second SQLite connection. The old
        # request transaction held its flushed audit INSERT open, blocking this
        # UPDATE until busy_timeout; an in-memory DB or mocked commit misses it.
        work_database = Database(operator.database.url)
        try:
            async with work_database.session() as work_session:
                await work_session.execute(text("PRAGMA busy_timeout=200"))
                await work_session.execute(
                    update(Site)
                    .where(Site.id == own_site_id)
                    .values(name=f"Completed delegated {action}")
                )
                support_events = (
                    await work_session.scalars(
                        select(AuditEvent).where(
                            AuditEvent.actor_user_id == operator.user_id,
                            AuditEvent.organization_id == organization_id,
                            AuditEvent.action == "support.access_request",
                        )
                    )
                ).all()
                matching_events = [
                    event
                    for event in support_events
                    if event.details.get("method") == "POST" and event.details.get("path") == path
                ]
                assert len(matching_events) == 1, (
                    "Support access audit must be durable before delegated work starts"
                )
                await work_session.commit()
        finally:
            await work_database.dispose()

        if action == "analyze-preview":
            return Analysis(
                significant=True,
                headline="Preview succeeded",
                summary="The separate database session completed.",
                cost=CostEstimate("scripted-support-test", TokenUsage(), 0.0),
            )
        return RunResult(
            site_id=own_site_id,
            change_id=own_change_id,
            status=CheckStatus.CHANGED,
        )

    delegate_methods = {
        "check": (operator.app.state.scheduler, "run_manual_check"),
        "snapshot": (operator.app.state.scheduler, "run_manual_snapshot"),
        "retry": (operator.app.state.runner, "reprocess_change"),
        "analyze": (operator.app.state.runner, "analyze_only"),
        "analyze-preview": (operator.app.state.runner, "analyze_preview"),
    }
    delegate_owner, delegate_method = delegate_methods[action]
    monkeypatch.setattr(delegate_owner, delegate_method, delegated_work)

    denied = await operator.client.post(path, headers=headers, json=payload)
    assert denied.status_code == 403, denied.text
    assert calls == []
    await _grant(operator, organization_id)
    wrong_organization = await operator.client.post(other_path, headers=headers, json=payload)
    assert wrong_organization.status_code == 404, wrong_organization.text
    assert calls == []

    completed = await operator.client.post(path, headers=headers, json=payload)
    assert completed.status_code == 200, completed.text
    assert calls == [resource_id]
    async with operator.database.session() as database_session:
        persisted_site = await database_session.get(Site, own_site_id)
        assert persisted_site is not None
        assert persisted_site.name == f"Completed delegated {action}"


async def test_public_acting_org_is_locked_until_audited_grant(
    public_operator_session: PublicOperatorSession,
) -> None:
    operator = public_operator_session
    client = operator.client
    organization_id = await _create_organization(operator, "Read only customer")
    headers = {"X-Acting-Org": str(organization_id)}

    instance_audit = await client.get("/api/audit-events")
    assert instance_audit.status_code == 200, instance_audit.text
    assert all(event["organization_id"] is None for event in instance_audit.json())
    assert "Read only customer" not in instance_audit.text

    denied_audit = await client.get("/api/audit-events", headers=headers)
    assert denied_audit.status_code == 403
    assert denied_audit.json()["error_code"] == "permission_denied"

    denied_read = await client.get("/api/sites", headers=headers)
    assert denied_read.status_code == 403
    assert denied_read.json()["error_code"] == "permission_denied"

    denied = await client.post(
        "/api/sites",
        headers=headers,
        json={"url": "https://customer.example.test"},
    )
    assert denied.status_code == 403
    assert denied.json()["error_code"] == "permission_denied"

    no_step_up = await client.post(
        "/api/support-access",
        headers=headers,
        json={"reason": "Investigating customer incident"},
    )
    assert no_step_up.status_code == 428

    granted = await _grant(operator, organization_id)
    assert granted["required"] is True
    assert granted["access_enabled"] is True
    assert granted["organization_id"] == organization_id
    assert granted["expires_at"] is not None

    created = await client.post(
        "/api/sites",
        headers=headers,
        json={"url": "https://customer.example.test"},
    )
    assert created.status_code == 201, created.text
    listed = await client.get("/api/sites", headers=headers)
    assert listed.status_code == 200, listed.text
    assert [site["id"] for site in listed.json()] == [created.json()["id"]]

    audit = await client.get("/api/audit-events", headers=headers, params={"limit": 20})
    actions = [event["action"] for event in audit.json()]
    assert "support.access_granted" in actions
    assert "support.access_request" in actions
    grant_event = next(
        event for event in audit.json() if event["action"] == "support.access_granted"
    )
    assert grant_event["details"]["reason"] == "Investigating customer incident"
    assert grant_event["details"]["ticket"] == "SUP-123"
    assert _OPERATOR_PASSWORD not in audit.text


async def test_omitting_or_switching_tenant_cannot_bypass_support_scope(
    public_operator_session: PublicOperatorSession,
) -> None:
    operator = public_operator_session
    client = operator.client
    organization_a = await _create_organization(operator, "Customer A")
    organization_b = await _create_organization(operator, "Customer B")
    await _grant(operator, organization_a)

    missing_scope = await client.post(
        "/api/sites",
        json={"url": "https://missing-scope.example.test"},
    )
    assert missing_scope.status_code == 400

    wrong_scope = await client.post(
        "/api/sites",
        headers={"X-Acting-Org": str(organization_b)},
        json={"url": "https://wrong-scope.example.test"},
    )
    assert wrong_scope.status_code == 403


async def test_instance_scope_cannot_read_tenant_notification_export_or_usage_data(
    public_operator_session: PublicOperatorSession,
) -> None:
    operator = public_operator_session
    client = operator.client
    organization_a = await _create_organization(operator, "Customer privacy A")
    organization_b = await _create_organization(operator, "Customer privacy B")
    await _seed_tenant_history(operator, organization_a, marker="tenant-a")
    await _seed_tenant_history(operator, organization_b, marker="tenant-b")

    tenant_paths = (
        "/api/notifications",
        "/api/exports/changes.xlsx",
        "/api/usage/summary",
        "/api/usage/verdicts",
    )
    for path in tenant_paths:
        response = await client.get(path)
        assert response.status_code == 400, (path, response.text)

    headers = {"X-Acting-Org": str(organization_a)}
    for path in tenant_paths:
        response = await client.get(path, headers=headers)
        assert response.status_code == 403, (path, response.text)

    await _grant(operator, organization_a)

    notifications = await client.get("/api/notifications", headers=headers)
    assert notifications.status_code == 200, notifications.text
    assert [row["recipient_email"] for row in notifications.json()] == ["tenant-a@example.test"]

    usage = await client.get("/api/usage/summary", headers=headers)
    assert usage.status_code == 200, usage.text
    assert usage.json()["calls"] == 1
    assert [row["model"] for row in usage.json()["by_model"]] == ["model-tenant-a"]

    verdicts = await client.get("/api/usage/verdicts", headers=headers)
    assert verdicts.status_code == 200, verdicts.text
    assert [row["label"] for row in verdicts.json()["by_site"]] == ["Site tenant-a"]

    exported = await client.get("/api/exports/changes.xlsx", headers=headers)
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][2] == "https://tenant-a.example.test"


async def test_workspace_branding_follows_the_authorized_acting_organization(
    public_operator_session: PublicOperatorSession,
) -> None:
    operator = public_operator_session
    client = operator.client
    organization_a = await _create_organization(operator, "Branded customer A")
    organization_b = await _create_organization(operator, "Branded customer B")
    headers_a = {"X-Acting-Org": str(organization_a)}
    headers_b = {"X-Acting-Org": str(organization_b)}

    await _grant(operator, organization_a)
    saved = await client.put(
        "/api/settings",
        headers=headers_a,
        json={"brand_name": "Customer A workspace"},
    )
    assert saved.status_code == 200, saved.text

    workspace_a = await client.get("/api/branding/workspace", headers=headers_a)
    assert workspace_a.status_code == 200, workspace_a.text
    assert workspace_a.json()["brand_name"] == "Customer A workspace"
    assert "no-store" in workspace_a.headers["cache-control"]

    # The cacheable public endpoint never varies by session or acting-org header.
    public_brand = await client.get("/api/branding", headers=headers_a)
    assert public_brand.status_code == 200
    assert public_brand.json()["brand_name"] == ""
    assert public_brand.headers["cache-control"] == "public, max-age=60"

    denied_switch = await client.get("/api/branding/workspace", headers=headers_b)
    assert denied_switch.status_code == 403

    await _grant(operator, organization_b)
    workspace_b = await client.get("/api/branding/workspace", headers=headers_b)
    assert workspace_b.status_code == 200, workspace_b.text
    assert workspace_b.json()["brand_name"] == ""


async def test_revocation_and_token_version_invalidate_write_access(
    public_operator_session: PublicOperatorSession,
) -> None:
    operator = public_operator_session
    client = operator.client
    organization_id = await _create_organization(operator, "Revoked customer")
    headers = {"X-Acting-Org": str(organization_id)}
    await _grant(operator, organization_id)
    session_cookie = client.cookies.get("driftwatch_session")
    support_cookie = client.cookies.get("driftwatch_support_access")
    assert session_cookie is not None
    assert support_cookie is not None
    operator.replay_client.cookies.set(
        "driftwatch_session",
        session_cookie,
        domain="driftwatch.example.test",
        path="/api",
    )
    operator.replay_client.cookies.set(
        "driftwatch_support_access",
        support_cookie,
        domain="driftwatch.example.test",
        path="/api",
    )

    revoked = await client.delete("/api/support-access", headers=headers)
    assert revoked.status_code == 204
    denied = await operator.replay_client.post(
        "/api/sites",
        headers=headers,
        json={"url": "https://revoked.example.test"},
    )
    assert denied.status_code == 403
    status_response = await operator.replay_client.get("/api/support-access", headers=headers)
    assert status_response.status_code == 200
    assert status_response.json()["access_enabled"] is False
    async with operator.database.session() as database_session:
        grant = await database_session.scalar(
            select(SupportAccessGrant).where(
                SupportAccessGrant.organization_id == organization_id,
                SupportAccessGrant.grant_event_id.is_not(None),
            )
        )
        assert grant is not None
        assert grant.revoked_at is not None

    await _grant(operator, organization_id)
    wrong_version = issue_support_access_token(
        operator.user_id,
        organization_id=organization_id,
        grant_event_id=1,
        secret="session-key-for-public-support-tests-0123456789",
        ttl_seconds=60,
        token_version=999,
        session_generation=_TEST_SESSION_GENERATION,
    )
    client.cookies.set(
        "driftwatch_support_access",
        wrong_version,
        domain="driftwatch.example.test",
        path="/api",
    )
    invalidated = await client.post(
        "/api/sites",
        headers=headers,
        json={"url": "https://invalidated.example.test"},
    )
    assert invalidated.status_code == 403


async def test_expired_or_malformed_grants_and_blank_reasons_are_rejected(
    public_operator_session: PublicOperatorSession,
) -> None:
    operator = public_operator_session
    client = operator.client
    organization_id = await _create_organization(operator, "Expired customer")
    headers = {"X-Acting-Org": str(organization_id)}
    await _step_up(operator, organization_id)

    blank = await client.post(
        "/api/support-access",
        headers=headers,
        json={"reason": "              "},
    )
    assert blank.status_code == 422

    expired = issue_support_access_token(
        operator.user_id,
        organization_id=organization_id,
        grant_event_id=1,
        secret="session-key-for-public-support-tests-0123456789",
        ttl_seconds=-1,
        session_generation=_TEST_SESSION_GENERATION,
    )
    client.cookies.set(
        "driftwatch_support_access",
        expired,
        domain="driftwatch.example.test",
        path="/api",
    )
    denied = await client.post(
        "/api/sites",
        headers=headers,
        json={"url": "https://expired.example.test"},
    )
    assert denied.status_code == 403


async def test_database_expiry_is_authoritative_for_status_and_writes(
    public_operator_session: PublicOperatorSession,
) -> None:
    operator = public_operator_session
    organization_id = await _create_organization(operator, "Database expired customer")
    headers = {"X-Acting-Org": str(organization_id)}
    await _grant(operator, organization_id)

    async with operator.database.session() as database_session:
        grant = await database_session.scalar(
            select(SupportAccessGrant).where(
                SupportAccessGrant.organization_id == organization_id,
                SupportAccessGrant.operator_user_id == operator.user_id,
            )
        )
        assert grant is not None
        grant.expires_at = datetime.now(UTC) - timedelta(seconds=1)
        await database_session.commit()

    status_response = await operator.client.get("/api/support-access", headers=headers)
    assert status_response.status_code == 200
    assert status_response.json()["access_enabled"] is False
    denied = await operator.client.post(
        "/api/sites",
        headers=headers,
        json={"url": "https://database-expired.example.test"},
    )
    assert denied.status_code == 403


async def test_rotated_and_cookie_less_revocation_invalidate_every_copied_grant(
    public_operator_session: PublicOperatorSession,
) -> None:
    operator = public_operator_session
    client = operator.client
    organization_id = await _create_organization(operator, "Rotated support customer")
    headers = {"X-Acting-Org": str(organization_id)}

    await _grant(operator, organization_id)
    old_cookie = client.cookies.get("driftwatch_support_access")
    session_cookie = client.cookies.get("driftwatch_session")
    assert old_cookie is not None
    assert session_cookie is not None
    await _grant(operator, organization_id)
    new_cookie = client.cookies.get("driftwatch_support_access")
    assert new_cookie is not None and new_cookie != old_cookie

    operator.replay_client.cookies.set(
        "driftwatch_session",
        session_cookie,
        domain="driftwatch.example.test",
        path="/api",
    )
    operator.replay_client.cookies.set(
        "driftwatch_support_access",
        old_cookie,
        domain="driftwatch.example.test",
        path="/api",
    )
    old_replay = await operator.replay_client.post(
        "/api/sites",
        headers=headers,
        json={"url": "https://old-grant.example.test"},
    )
    assert old_replay.status_code == 403

    operator.replay_client.cookies.set(
        "driftwatch_support_access",
        new_cookie,
        domain="driftwatch.example.test",
        path="/api",
    )
    client.cookies.delete(
        "driftwatch_support_access",
        domain="driftwatch.example.test",
        path="/api",
    )
    revoked = await client.delete("/api/support-access", headers=headers)
    assert revoked.status_code == 204
    new_replay = await operator.replay_client.post(
        "/api/sites",
        headers=headers,
        json={"url": "https://new-grant.example.test"},
    )
    assert new_replay.status_code == 403

    async with operator.database.session() as database_session:
        grants = (
            (
                await database_session.execute(
                    select(SupportAccessGrant).where(
                        SupportAccessGrant.organization_id == organization_id,
                        SupportAccessGrant.operator_user_id == operator.user_id,
                    )
                )
            )
            .scalars()
            .all()
        )
    assert len(grants) == 2
    assert all(grant.revoked_at is not None for grant in grants)
    assert all(grant.active_marker is None for grant in grants)


async def test_public_dead_check_redrive_requires_support_grant_and_step_up(
    public_operator_session: PublicOperatorSession,
) -> None:
    operator = public_operator_session
    organization_id = await _create_organization(operator, "Recovery support customer")
    now = datetime.now(UTC)
    async with operator.database.session() as database_session:
        site = Site(
            organization_id=organization_id,
            url="https://recovery-customer.example.test",
        )
        database_session.add(site)
        await database_session.flush()
        dead = SiteCheckJob(
            organization_id=organization_id,
            site_id=site.id,
            idempotency_key="support:redrive:dead",
            kind=CheckJobKind.CHECK,
            source=CheckJobSource.SCHEDULED,
            analyze=True,
            status=CheckJobStatus.DEAD,
            attempt_count=5,
            available_at=None,
            enqueued_at=now,
            completed_at=now,
            last_error="customer-private failure",
        )
        database_session.add(dead)
        await database_session.commit()
        dead_id = dead.id

    instance = await operator.client.get("/api/operations/check-jobs/dead")
    assert instance.status_code == 200, instance.text
    assert [row["id"] for row in instance.json()["items"]] == [dead_id]
    assert "customer-private" not in instance.text
    assert "recovery-customer" not in instance.text

    headers = {"X-Acting-Org": str(organization_id)}
    denied = await operator.client.get("/api/operations/tenant/check-jobs/dead", headers=headers)
    assert denied.status_code == 403

    await _grant(operator, organization_id)
    tenant = await operator.client.get("/api/operations/tenant/check-jobs/dead", headers=headers)
    assert tenant.status_code == 200, tenant.text
    assert [row["id"] for row in tenant.json()["items"]] == [dead_id]

    operator.client.cookies.delete(STEP_UP_COOKIE)
    no_step_up = await operator.client.post(
        f"/api/operations/tenant/check-jobs/dead/{dead_id}/redrive",
        headers=headers,
        json={
            "idempotency_key": "public-redrive-001",
            "reason": "The customer confirmed the upstream incident is resolved",
            "ticket": "SUP-991",
        },
    )
    assert no_step_up.status_code == 428

    await _step_up(operator, organization_id)
    recovered = await operator.client.post(
        f"/api/operations/tenant/check-jobs/dead/{dead_id}/redrive",
        headers=headers,
        json={
            "idempotency_key": "public-redrive-001",
            "reason": "The customer confirmed the upstream incident is resolved",
            "ticket": "SUP-991",
        },
    )
    assert recovered.status_code == 200, recovered.text
    assert recovered.json()["created"] is True
